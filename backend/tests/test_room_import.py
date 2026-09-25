"""Bulk room import and export from Room Management."""
import io
import uuid

import openpyxl

from app.models.inpatient import Bed, RoomManagement


def _xlsx(rooms, beds=None):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Rooms"
    sheet.append([
        "room_number", "room_type", "floor", "department", "ward",
        "bed_count", "room_charge_per_day", "nursing_charge_per_visit",
        "amenities", "is_isolation", "gender_policy",
    ])
    for row in rooms:
        sheet.append(row)
    beds_sheet = wb.create_sheet("Beds")
    beds_sheet.append(["room_number", "bed_label"])
    for row in beds or []:
        beds_sheet.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, auth_headers, content, filename="rooms.xlsx"):
    return client.post(
        "/api/inpatient/rooms/import",
        headers=auth_headers,
        files={"file": (
            filename,
            io.BytesIO(content),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )


def _deactivate(db_session, numbers):
    rooms = db_session.query(RoomManagement).filter(
        RoomManagement.room_number.in_(numbers)
    ).all()
    for room in rooms:
        db_session.query(Bed).filter(Bed.room_id == room.id).delete()
        room.is_active = False
    db_session.commit()


def test_import_creates_rooms_and_export_round_trips(client, auth_headers, db_session):
    suffix = uuid.uuid4().hex[:6].upper()
    general = f"IMP-{suffix}-G"
    icu = f"IMP-{suffix}-I"
    try:
        created = _upload(client, auth_headers, _xlsx(
            [
                [general, "general", "1", "Medicine", "Male Ward", 2, 1500, 100, "ac;wifi", "false", "male"],
                [icu, "icu", "2", "Critical Care", "ICU", 1, 8000, 350, "oxygen_point", "true", "mixed"],
            ],
            [[general, "A"], [general, "B"]],
        ))
        assert created.status_code == 200, created.text
        body = created.json()
        assert body["ok"] is True
        assert body["created_rooms"] == 2
        assert body["created_beds"] == 3

        listed = client.get("/api/inpatient/rooms", headers=auth_headers).json()
        by_number = {room["room_number"]: room for room in listed}
        assert by_number[general]["room_type"] == "general"
        assert by_number[general]["bed_count"] == 2
        assert by_number[icu]["is_isolation"] is True

        exported = client.get("/api/inpatient/rooms/export", headers=auth_headers)
        assert exported.status_code == 200, exported.text
        wb = openpyxl.load_workbook(io.BytesIO(exported.content))
        numbers = {row[0] for row in wb["Rooms"].iter_rows(min_row=2, values_only=True)}
        assert general in numbers
        assert icu in numbers

        again = _upload(client, auth_headers, exported.content)
        assert again.status_code == 200, again.text
        assert again.json()["ok"] is True
        assert again.json()["created_rooms"] == 0
        assert again.json()["skipped"] >= 2
    finally:
        _deactivate(db_session, [general, icu])


def test_import_creates_a_new_room_type(client, auth_headers, db_session):
    suffix = uuid.uuid4().hex[:6].upper()
    number = f"PH-{suffix}"
    try:
        resp = _upload(client, auth_headers, _xlsx([
            [number, "Penthouse Suite", "3", "Luxury", "Deluxe", 1, 12000, 0, "", "false", "mixed"],
        ]))
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["ok"] is True
        assert body["created_rooms"] == 1
        listed = client.get("/api/inpatient/rooms", headers=auth_headers).json()
        room = next(row for row in listed if row["room_number"] == number)
        assert room["room_type"] == "penthouse_suite"
        types = client.get("/api/inpatient/room-types", headers=auth_headers).json()
        match = next(row for row in types if row["value"] == "penthouse_suite")
        assert match["label"] == "Penthouse Suite"
    finally:
        _deactivate(db_session, [number])
