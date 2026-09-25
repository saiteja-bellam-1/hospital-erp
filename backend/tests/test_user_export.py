"""Excel export of the user roster. Passwords must never appear in the file."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.user import User, UserRole
from app.utils.auth import create_access_token, get_password_hash


@pytest.fixture(scope="module")
def hospital_admin(TestSessionLocal, seed_data):
    session = TestSessionLocal()
    try:
        role = session.query(UserRole).filter_by(name="hospital_admin").first()
        if role is None:
            role = UserRole(name="hospital_admin", is_system_role=True)
            session.add(role)
            session.flush()
        user = session.query(User).filter_by(username="testhospadmin").first()
        if user is None:
            user = User(
                username="testhospadmin",
                password_hash=get_password_hash("hosp123"),
                email="hospadmin@test.com",
                first_name="Hosp",
                last_name="Admin",
                role_id=role.id,
                hospital_id=seed_data["hospital_id"],
                is_active=True,
            )
            session.add(user)
            session.commit()
        token = create_access_token(data={"sub": "testhospadmin"})
        return {"Authorization": f"Bearer {token}"}
    finally:
        session.close()


def _sheet(client, headers):
    res = client.get("/api/admin/users/export/xlsx", headers=headers)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(res.content))
    sheet = wb.active
    header = [cell.value for cell in sheet[1]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    return header, rows


def test_export_includes_roster_without_passwords(client, auth_headers):
    header, rows = _sheet(client, auth_headers)
    assert "Password" not in header
    assert header[:4] == ["Username", "Email", "First Name", "Last Name"]
    by_username = {row[0]: row for row in rows}
    assert "testadmin" in by_username
    assert "testdoctor" in by_username
    assert by_username["testdoctor"][2:4] == ("Dr", "Smith")
    assert by_username["testdoctor"][5] == "doctor"
    joined = " ".join(str(cell or "") for row in rows for cell in row).lower()
    assert "doctor123" not in joined
    assert "admin123" not in joined


def test_hospital_admin_export_hides_vendor_account(client, hospital_admin):
    header, rows = _sheet(client, hospital_admin)
    usernames = {row[0] for row in rows}
    assert "testadmin" not in usernames
    assert "testdoctor" in usernames
    assert "Password" not in header


def test_non_admin_cannot_export_users(client):
    token = create_access_token(data={"sub": "testreceptionist"})
    res = client.get(
        "/api/admin/users/export/xlsx",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert res.status_code == 403
