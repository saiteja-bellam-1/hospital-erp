"""Export and import of room-type nursing rates and doctor visit overrides."""
import io

import openpyxl

from app.models.inpatient import DoctorRoomTypeRate, RoomTypeRateConfig


def _cleanup(db_session, hospital_id):
    db_session.query(DoctorRoomTypeRate).filter(
        DoctorRoomTypeRate.hospital_id == hospital_id
    ).delete()
    db_session.query(RoomTypeRateConfig).filter(
        RoomTypeRateConfig.hospital_id == hospital_id
    ).delete()
    db_session.commit()


def _xlsx(nursing_rows, doctor_rows=None):
    wb = openpyxl.Workbook()
    nursing = wb.active
    nursing.title = "Nursing rates"
    nursing.append(["room_type", "room_type_label", "nursing_charge_per_visit"])
    for row in nursing_rows:
        nursing.append(row)
    if doctor_rows is not None:
        doctor = wb.create_sheet("Doctor room rates")
        doctor.append(["doctor_username", "doctor_name", "room_type", "visit_rate"])
        for row in doctor_rows:
            doctor.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, auth_headers, content, filename="room_type_configuration.xlsx"):
    media = (
        "text/csv" if filename.endswith(".csv")
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return client.post(
        "/api/inpatient/room-type-rates/import",
        headers=auth_headers,
        files={"file": (filename, io.BytesIO(content), media)},
    )


def test_export_lists_room_types_and_round_trips(client, auth_headers, db_session, seed_data):
    hospital_id = seed_data["hospital_id"]
    _cleanup(db_session, hospital_id)
    try:
        imported = _upload(client, auth_headers, _xlsx(
            [["icu", "ICU", 350], ["private", "Private", 200]],
            [["testdoctor", "Dr Smith", "icu", 1500]],
        ))
        assert imported.status_code == 200, imported.text
        body = imported.json()
        assert body["ok"] is True
        assert body["nursing_updated"] == 2
        assert body["doctor_rates_upserted"] == 1

        exported = client.get("/api/inpatient/room-type-rates/export", headers=auth_headers)
        assert exported.status_code == 200, exported.text
        wb = openpyxl.load_workbook(io.BytesIO(exported.content))
        assert "Nursing rates" in wb.sheetnames
        assert "Doctor room rates" in wb.sheetnames
        nursing = {
            row[0]: row[2]
            for row in wb["Nursing rates"].iter_rows(min_row=2, values_only=True)
        }
        assert nursing["icu"] == 350
        assert nursing["private"] == 200
        assert nursing["general"] is None
        doctor_rows = list(wb["Doctor room rates"].iter_rows(min_row=2, values_only=True))
        assert doctor_rows == [("testdoctor", "Dr Smith", "icu", 1500)]

        again = _upload(client, auth_headers, exported.content, "room_type_configuration.xlsx")
        assert again.status_code == 200, again.text
        assert again.json()["ok"] is True
        rates = client.get("/api/inpatient/room-type-rates", headers=auth_headers).json()
        by_type = {row["room_type"]: row["nursing_charge_per_visit"] for row in rates}
        assert by_type["icu"] == 350
        assert by_type["private"] == 200
    finally:
        _cleanup(db_session, hospital_id)


def test_import_rejects_invalid_rows_without_saving(client, auth_headers, db_session, seed_data):
    hospital_id = seed_data["hospital_id"]
    _cleanup(db_session, hospital_id)
    try:
        resp = _upload(client, auth_headers, _xlsx(
            [["icu", "ICU", 100], ["penthouse", "Penthouse", 900]],
            [["nobody", "", "icu", 10]],
        ))
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["ok"] is False
        assert body["nursing_updated"] == 0
        messages = " ".join(err["message"] for err in body["errors"])
        assert "nobody" in messages
        saved = db_session.query(RoomTypeRateConfig).filter(
            RoomTypeRateConfig.hospital_id == hospital_id
        ).count()
        assert saved == 0
    finally:
        _cleanup(db_session, hospital_id)


def test_csv_imports_nursing_rates_only(client, auth_headers, db_session, seed_data):
    hospital_id = seed_data["hospital_id"]
    _cleanup(db_session, hospital_id)
    try:
        csv_body = "room_type,nursing_charge_per_visit\nhdu,250\n"
        resp = _upload(client, auth_headers, csv_body.encode(), "nursing.csv")
        assert resp.status_code == 200, resp.text
        assert resp.json()["ok"] is True
        assert resp.json()["nursing_updated"] == 1
        rates = client.get("/api/inpatient/room-type-rates", headers=auth_headers).json()
        hdu = next(row for row in rates if row["room_type"] == "hdu")
        assert hdu["nursing_charge_per_visit"] == 250
    finally:
        _cleanup(db_session, hospital_id)
