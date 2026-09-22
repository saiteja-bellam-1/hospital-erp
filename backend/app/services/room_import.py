"""Inpatient room/bed catalog import and export.

Workbook layout matches Guided Setup (`Rooms` + `Beds` sheets) so a setup file
and a Room Management export are interchangeable. Occupancy is never imported.
"""
from __future__ import annotations

import io
import json
from typing import Any, Optional

import openpyxl
from sqlalchemy.orm import Session

from app.models.inpatient import Bed, RoomManagement
from app.services.onboarding_import import (
    AMENITY_OPTIONS,
    ROOM_TYPES,
    _cell_bool,
    _cell_float,
    _cell_int,
    _cell_str,
    _row_empty,
    read_csv_or_xlsx,
)

ROOM_HEADERS = [
    "room_number", "room_type", "floor", "department", "ward",
    "bed_count", "room_charge_per_day", "nursing_charge_per_visit",
    "amenities", "is_isolation", "gender_policy",
]
BED_HEADERS = ["room_number", "bed_label"]
GENDER_POLICIES = {"mixed", "male", "female"}


def _empty_summary(*, dry_run: bool) -> dict:
    return {
        "dry_run": dry_run,
        "total_rows": 0,
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "created_beds": 0,
        "error_count": 0,
        "errors": [],
        "preview": [],
    }


def _preview(*, row: int, key: str, name: str, sheet: str, status: str, message: str = "") -> dict:
    return {
        "row": row,
        "key": key,
        "name": name,
        "sheet": sheet,
        "status": status,
        "message": message,
    }


def _error(summary: dict, *, sheet: str, row: int, message: str, key: str = "", name: str = "") -> None:
    summary["errors"].append({"sheet": sheet, "row": row, "message": message})
    summary["preview"].append(_preview(
        row=row, key=key, name=name, sheet=sheet, status="error", message=message,
    ))


def _norm_room_type(value: Any) -> str:
    return (_cell_str(value) or "").strip().lower().replace(" ", "_")


def _parse_amenities(raw: Any) -> list[str]:
    text = _cell_str(raw) or ""
    if not text:
        return []
    return [a.strip() for a in text.replace(";", ",").split(",") if a.strip()]


def _stored_amenities(raw: Any) -> list[str]:
    if not raw:
        return []
    if isinstance(raw, list):
        return [str(a) for a in raw if str(a).strip()]
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return [str(a) for a in parsed if str(a).strip()]
    except (TypeError, ValueError):
        pass
    return _parse_amenities(raw)


def _export_amenities(raw: Any) -> str:
    return ";".join(_stored_amenities(raw))


def _amenities_json(amenities: list[str]) -> Optional[str]:
    return json.dumps(amenities) if amenities else None


def _bool_cell(value: Any) -> str:
    return "true" if value else "false"


def _index_rooms(db: Session, hospital_id: int) -> dict[str, RoomManagement]:
    rooms = db.query(RoomManagement).filter(RoomManagement.hospital_id == hospital_id).all()
    by_number: dict[str, RoomManagement] = {}
    for room in rooms:
        key = (room.room_number or "").lower()
        if not key:
            continue
        existing = by_number.get(key)
        if existing is None or (not existing.is_active and room.is_active):
            by_number[key] = room
    return by_number


def _sync_room_counts(db: Session, room: RoomManagement) -> None:
    beds = db.query(Bed).filter(Bed.room_id == room.id).all()
    room.bed_count = len(beds)
    room.available_beds = sum(1 for b in beds if (b.status or "") == "available")
    room.is_occupied = room.bed_count > 0 and room.available_beds == 0


def _bed_protected(bed: Bed) -> bool:
    if bed.current_admission_id:
        return True
    return (bed.status or "").lower() != "available"


def _next_auto_label(existing_lower: set[str], start_at: int = 1) -> str:
    n = max(1, start_at)
    while True:
        label = f"Bed-{n}"
        if label.lower() not in existing_lower:
            return label
        n += 1


def _add_bed(db: Session, room: RoomManagement, label: str, summary: dict) -> Bed:
    bed = Bed(room_id=room.id, bed_label=label, status="available")
    db.add(bed)
    summary["created_beds"] += 1
    return bed


def _auto_create_beds(db: Session, room: RoomManagement, count: int, summary: dict) -> None:
    existing = db.query(Bed).filter(Bed.room_id == room.id).all()
    labels = {(b.bed_label or "").lower() for b in existing}
    created = 0
    n = 1
    while len(existing) + created < count:
        label = _next_auto_label(labels, n)
        labels.add(label.lower())
        _add_bed(db, room, label, summary)
        created += 1
        n += 1


def _shrink_beds(db: Session, room: RoomManagement, target: int) -> Optional[str]:
    beds = (
        db.query(Bed)
        .filter(Bed.room_id == room.id)
        .order_by(Bed.id.desc())
        .all()
    )
    if len(beds) <= target:
        return None
    protected = [b for b in beds if _bed_protected(b)]
    if len(protected) > target:
        return (
            f"Cannot reduce beds below currently occupied/unavailable count "
            f"({len(protected)})"
        )
    to_remove = len(beds) - target
    removable = [b for b in beds if not _bed_protected(b)]
    for bed in removable[:to_remove]:
        db.delete(bed)
    return None


def _apply_catalog(
    room: RoomManagement,
    *,
    room_type: str,
    floor: Optional[str],
    department: Optional[str],
    ward: Optional[str],
    charge: float,
    nursing: float,
    amenities: list[str],
    is_isolation: bool,
    gender: str,
) -> None:
    room.room_type = room_type
    room.floor = floor
    room.department = department
    room.ward = ward
    room.room_charge_per_day = charge
    room.nursing_charge_per_visit = nursing
    room.amenities = _amenities_json(amenities)
    room.is_isolation = is_isolation
    room.gender_policy = gender
    room.is_active = True


def _parse_room_row(row: dict) -> tuple[dict, list[str]]:
    errors: list[str] = []
    room_number = _cell_str(row.get("room_number"))
    room_type = _norm_room_type(row.get("room_type"))
    try:
        charge = _cell_float(row.get("room_charge_per_day"), field="room_charge_per_day")
    except ValueError as exc:
        charge = None
        errors.append(str(exc))
    try:
        bed_count = _cell_int(row.get("bed_count"), field="bed_count", default=1)
    except ValueError as exc:
        bed_count = None
        errors.append(str(exc))
    try:
        nursing = _cell_float(row.get("nursing_charge_per_visit"), field="nursing_charge_per_visit")
    except ValueError as exc:
        nursing = None
        errors.append(str(exc))

    if not room_number:
        errors.append("room_number is required")
    elif len(room_number) > 20:
        errors.append("room_number must be 20 characters or fewer")
    if not room_type:
        errors.append("room_type is required")
    elif room_type not in ROOM_TYPES:
        errors.append(f"Invalid room_type '{room_type}'")
    if charge is None and "room_charge_per_day must be a number" not in errors:
        errors.append("room_charge_per_day is required")
    elif charge is not None and charge < 0:
        errors.append("room_charge_per_day cannot be negative")
    if bed_count is None:
        pass
    elif bed_count < 1:
        errors.append("bed_count must be at least 1")
    nursing = nursing or 0.0
    if nursing < 0:
        errors.append("nursing_charge_per_visit cannot be negative")

    gender = (_cell_str(row.get("gender_policy")) or "mixed").lower()
    if gender not in GENDER_POLICIES:
        errors.append("gender_policy must be mixed, male or female")

    amenities = _parse_amenities(row.get("amenities"))
    bad = [a for a in amenities if a not in AMENITY_OPTIONS]
    if bad:
        errors.append(f"Unknown amenities: {', '.join(bad)}")

    parsed = {
        "room_number": room_number or "",
        "room_type": room_type,
        "floor": _cell_str(row.get("floor")),
        "department": _cell_str(row.get("department")),
        "ward": _cell_str(row.get("ward")),
        "bed_count": bed_count or 1,
        "charge": charge,
        "nursing": nursing,
        "amenities": amenities,
        "is_isolation": _cell_bool(row.get("is_isolation"), default=False),
        "gender": gender,
    }
    return parsed, errors


def _workbook_bytes(build_fn) -> bytes:
    wb = openpyxl.Workbook()
    build_fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_rooms_template() -> bytes:
    def build(wb: openpyxl.Workbook) -> None:
        ws = wb.active
        ws.title = "Rooms"
        ws.append(ROOM_HEADERS)
        ws.append([
            "G-101", "general", "1", "Medicine", "Male Ward",
            4, 1500, 100, "call_bell", "false", "male",
        ])
        ws.append([
            "ICU-1", "icu", "2", "Critical Care", "ICU",
            1, 8000, 350, "ac;cardiac_monitor;oxygen_point", "false", "mixed",
        ])

        beds = wb.create_sheet("Beds")
        beds.append(BED_HEADERS)
        beds.append(["G-101", "A"])
        beds.append(["G-101", "B"])
        beds.append(["G-101", "C"])
        beds.append(["G-101", "D"])
        beds.append(["ICU-1", "Bed-1"])

        notes = wb.create_sheet("Instructions")
        for line in [
            "KT HEALTH ERP — Room Import Template",
            "",
            "Fill the Rooms sheet — one row per room.",
            "  Required: room_number, room_type, room_charge_per_day.",
            "  Match key: room_number (case-insensitive, per hospital).",
            "  on_duplicate=skip (default) leaves existing active rooms alone.",
            "  on_duplicate=update updates catalog fields or reactivates a deactivated room.",
            "",
            "room_type must be one of: " + ", ".join(sorted(ROOM_TYPES)) + ".",
            "gender_policy: mixed, male or female.",
            "amenities: comma- or semicolon-separated keys such as ac, tv, wifi, oxygen_point.",
            "is_isolation: true/false, yes/no, 1/0.",
            "",
            "Beds sheet is optional. Columns: room_number, bed_label.",
            "If a room has no Beds rows, beds are created as Bed-1..Bed-N from bed_count.",
            "Occupancy (available_beds, bed status, admissions) is never imported.",
            "CSV files import the Rooms sheet only; beds are auto-created from bed_count.",
        ]:
            notes.append([line])

    return _workbook_bytes(build)


def export_rooms_xlsx(db: Session, hospital_id: int) -> bytes:
    rooms = (
        db.query(RoomManagement)
        .filter(
            RoomManagement.hospital_id == hospital_id,
            RoomManagement.is_active.is_(True),
        )
        .order_by(RoomManagement.room_number)
        .all()
    )

    def build(wb: openpyxl.Workbook) -> None:
        ws = wb.active
        ws.title = "Rooms"
        ws.append(ROOM_HEADERS)
        beds_sheet = wb.create_sheet("Beds")
        beds_sheet.append(BED_HEADERS)

        for room in rooms:
            ws.append([
                room.room_number or "",
                room.room_type or "",
                room.floor or "",
                room.department or "",
                room.ward or "",
                room.bed_count or 0,
                room.room_charge_per_day if room.room_charge_per_day is not None else "",
                float(room.nursing_charge_per_visit or 0),
                _export_amenities(room.amenities),
                _bool_cell(bool(room.is_isolation)),
                room.gender_policy or "mixed",
            ])
            beds = (
                db.query(Bed)
                .filter(Bed.room_id == room.id)
                .order_by(Bed.bed_label)
                .all()
            )
            for bed in beds:
                beds_sheet.append([room.room_number or "", bed.bed_label or ""])

    return _workbook_bytes(build)


def import_rooms_workbook(
    db: Session,
    hospital_id: int,
    content: bytes,
    filename: str,
    *,
    dry_run: bool,
    on_duplicate: str,
) -> dict:
    name = (filename or "").lower()
    if not (name.endswith(".xlsx") or name.endswith(".csv")):
        raise ValueError("Unsupported file type. Upload a .xlsx or .csv file.")
    if not content:
        raise ValueError("Uploaded file is empty")

    on_duplicate = on_duplicate if on_duplicate in ("skip", "update") else "skip"
    sheets = read_csv_or_xlsx(content, filename, ["Rooms", "Beds"])
    room_rows = sheets.get("Rooms") or []
    bed_rows = sheets.get("Beds") or []
    summary = _empty_summary(dry_run=dry_run)
    room_by_number = _index_rooms(db, hospital_id)

    mutated: set[int] = set()
    new_room_ids: set[int] = set()
    target_by_room: dict[int, tuple[int, int, str, str]] = {}

    for row in room_rows:
        if _row_empty(row):
            continue
        line = int(row.get("_row") or 0)
        summary["total_rows"] += 1
        parsed, row_errs = _parse_room_row(row)
        key = parsed["room_number"]
        if row_errs:
            _error(
                summary, sheet="Rooms", row=line, message="; ".join(row_errs),
                key=key, name=parsed["room_type"],
            )
            continue

        existing = room_by_number.get(key.lower())
        if existing and existing.is_active and on_duplicate == "skip":
            summary["skipped"] += 1
            summary["preview"].append(_preview(
                row=line, key=key, name=parsed["room_type"], sheet="Rooms",
                status="skip", message="Room number already exists",
            ))
            continue

        catalog_kwargs = dict(
            room_type=parsed["room_type"],
            floor=parsed["floor"],
            department=parsed["department"],
            ward=parsed["ward"],
            charge=parsed["charge"],
            nursing=parsed["nursing"],
            amenities=parsed["amenities"],
            is_isolation=parsed["is_isolation"],
            gender=parsed["gender"],
        )

        if existing:
            _apply_catalog(existing, **catalog_kwargs)
            mutated.add(existing.id)
            target_by_room[existing.id] = (parsed["bed_count"], line, key, parsed["room_type"])
            summary["updated"] += 1
            summary["preview"].append(_preview(
                row=line, key=key, name=parsed["room_type"], sheet="Rooms",
                status="update",
            ))
            continue

        room = RoomManagement(
            room_number=key,
            hospital_id=hospital_id,
            bed_count=parsed["bed_count"],
            available_beds=parsed["bed_count"],
            is_occupied=False,
            is_active=True,
        )
        _apply_catalog(room, **catalog_kwargs)
        db.add(room)
        db.flush()
        room_by_number[key.lower()] = room
        mutated.add(room.id)
        new_room_ids.add(room.id)
        target_by_room[room.id] = (parsed["bed_count"], line, key, parsed["room_type"])
        summary["created"] += 1
        summary["preview"].append(_preview(
            row=line, key=key, name=parsed["room_type"], sheet="Rooms",
            status="new",
        ))

    for row in bed_rows:
        if _row_empty(row):
            continue
        line = int(row.get("_row") or 0)
        room_number = _cell_str(row.get("room_number"))
        bed_label = _cell_str(row.get("bed_label"))
        preview_key = f"{room_number or ''}/{bed_label or ''}"
        if not room_number or not bed_label:
            _error(
                summary, sheet="Beds", row=line,
                message="room_number and bed_label are required",
                key=preview_key,
            )
            continue
        room = room_by_number.get(room_number.lower())
        if not room:
            _error(
                summary, sheet="Beds", row=line,
                message=f"Room '{room_number}' not found",
                key=preview_key, name=bed_label,
            )
            continue
        if room.id not in mutated:
            summary["skipped"] += 1
            summary["preview"].append(_preview(
                row=line, key=preview_key, name=bed_label, sheet="Beds",
                status="skip", message="Room was not created or updated",
            ))
            continue
        exists = (
            db.query(Bed)
            .filter(Bed.room_id == room.id, Bed.bed_label == bed_label)
            .first()
        )
        if exists:
            summary["skipped"] += 1
            summary["preview"].append(_preview(
                row=line, key=preview_key, name=bed_label, sheet="Beds",
                status="skip", message="Bed label already exists",
            ))
            continue
        _add_bed(db, room, bed_label, summary)
        summary["preview"].append(_preview(
            row=line, key=preview_key, name=bed_label, sheet="Beds",
            status="new",
        ))

    db.flush()

    for room_id, (target, line, key, room_type) in target_by_room.items():
        room = db.query(RoomManagement).filter(RoomManagement.id == room_id).first()
        if not room:
            continue
        current = db.query(Bed).filter(Bed.room_id == room.id).count()
        if room_id in new_room_ids:
            if current == 0:
                _auto_create_beds(db, room, target, summary)
        elif target > current:
            _auto_create_beds(db, room, target, summary)
        elif target < current:
            shrink_err = _shrink_beds(db, room, target)
            if shrink_err:
                _error(
                    summary, sheet="Rooms", row=line, message=shrink_err,
                    key=key, name=room_type,
                )
        db.flush()
        _sync_room_counts(db, room)

    summary["error_count"] = len(summary["errors"])
    return summary
