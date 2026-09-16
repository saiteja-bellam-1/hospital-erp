"""Build a single-sheet Excel workbook from EHR patient chart history."""
from __future__ import annotations

import io
import json
from typing import Any, Iterable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.export_branding import apply_workbook_branding, hospital_brand_dict

VALID_SECTIONS = (
    "timeline",
    "visits",
    "consultations",
    "prescriptions",
    "lab",
    "pharmacy",
    "billing",
    "documents",
)

_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_HEADER_FONT = Font(bold=True, size=10, color="1A202C")
_SECTION_FONT = Font(bold=True, size=11, color="1A365D")


def parse_sections(raw: Optional[str]) -> list[str]:
    """Parse comma-separated section keys; default to all tabs when empty."""
    if not raw or not str(raw).strip():
        return list(VALID_SECTIONS)
    seen = set()
    out: list[str] = []
    for part in str(raw).split(","):
        key = part.strip().lower()
        if key in VALID_SECTIONS and key not in seen:
            seen.add(key)
            out.append(key)
    return out


def build_patient_chart_xlsx(db, hospital_id: int, history: dict, sections: Sequence[str]) -> bytes:
    """Return .xlsx bytes with all selected chart tabs on one sheet."""
    hospital = hospital_brand_dict(db, hospital_id)
    patient = history.get("patient") or {}
    selected = [s for s in sections if s in VALID_SECTIONS] or list(VALID_SECTIONS)

    wb = Workbook()
    ws = wb.active
    ws.title = "Patient Chart"

    row = apply_workbook_branding(ws, hospital, _patient_meta(patient))
    row = _write_patient_header(ws, row, patient, history)
    row += 1

    writers = {
        "timeline": _write_timeline_section,
        "visits": _write_visits_section,
        "consultations": _write_consultations_section,
        "prescriptions": _write_prescriptions_section,
        "lab": _write_lab_section,
        "pharmacy": _write_pharmacy_section,
        "billing": _write_billing_section,
        "documents": _write_documents_section,
    }
    for key in selected:
        row = writers[key](ws, row, history)
        row += 1

    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _patient_meta(patient: dict) -> list[tuple[str, Any]]:
    return [
        ("Export", "Patient chart"),
        ("Patient", patient.get("full_name") or ""),
        ("Patient ID", patient.get("patient_id") or ""),
        ("Phone", patient.get("primary_phone") or ""),
        ("Age / Gender", f"{patient.get('age') or '—'} / {patient.get('gender') or '—'}"),
        ("Blood group", patient.get("blood_group") or "—"),
    ]


def _write_patient_header(ws, row: int, patient: dict, history: dict) -> int:
    row = _write_kv_block(ws, row, "Demographics", [
        ("Full name", patient.get("full_name")),
        ("Patient ID", patient.get("patient_id")),
        ("Date of birth", patient.get("date_of_birth")),
        ("Age", patient.get("age")),
        ("Gender", patient.get("gender")),
        ("Blood group", patient.get("blood_group")),
        ("Phone", patient.get("primary_phone")),
        ("Emergency phone", patient.get("emergency_contact_phone")),
        ("Address", patient.get("address")),
    ])
    row += 1

    allergies = history.get("allergies") or []
    row = _write_table(
        ws, row, "Allergies",
        ["Allergen", "Type", "Severity", "Reaction", "Active", "Recorded"],
        [
            [
                a.get("allergen"),
                a.get("allergy_type"),
                a.get("severity"),
                a.get("reaction"),
                "Yes" if a.get("is_active") else "No",
                a.get("recorded_at"),
            ]
            for a in allergies
        ],
    )
    row += 1

    med_hist = history.get("medical_history") or []
    return _write_table(
        ws, row, "Medical history",
        ["Condition", "Status", "Diagnosed", "Notes"],
        [
            [mh.get("condition"), mh.get("status"), mh.get("diagnosed_date"), mh.get("notes")]
            for mh in med_hist
        ],
    )


def _write_timeline_section(ws, row: int, history: dict) -> int:
    type_labels = {
        "consultation": "Consultation",
        "prescription": "Prescription",
        "lab_order": "Lab Order",
        "appointment": "Appointment",
        "admission": "Admission",
        "pharmacy_sale": "Pharmacy",
        "physio_session": "Physio Session",
    }
    rows = []
    for item in history.get("timeline") or []:
        data = item.get("data") or {}
        t = item.get("type")
        summary = ""
        if t == "consultation":
            summary = data.get("chief_complaint") or data.get("consultation_number") or ""
        elif t == "prescription":
            summary = data.get("diagnosis") or data.get("prescription_id") or ""
        elif t == "lab_order":
            summary = data.get("test_name") or data.get("order_number") or ""
        elif t == "appointment":
            summary = f"{data.get('doctor_name') or ''} · {data.get('appointment_number') or ''}".strip(" ·")
        elif t == "admission":
            summary = f"{data.get('admission_number') or ''} · {data.get('admission_type') or ''}".strip(" ·")
        elif t == "pharmacy_sale":
            summary = data.get("sale_number") or ""
        elif t == "physio_session":
            summary = data.get("service_name") or data.get("session_type") or data.get("appointment_number") or ""
        rows.append([
            item.get("date"),
            type_labels.get(t, t),
            summary,
            data.get("status") or data.get("payment_status") or "",
            data.get("doctor_name") or data.get("therapist_name") or "",
        ])
    return _write_table(ws, row, "Timeline", ["Date", "Type", "Summary", "Status", "Clinician"], rows)


def _write_visits_section(ws, row: int, history: dict) -> int:
    appointments = history.get("appointments") or []
    row = _write_table(
        ws, row, "Appointments",
        ["Date", "Number", "Doctor", "Status", "Payment", "Amount"],
        [
            [
                ap.get("appointment_date"),
                ap.get("appointment_number"),
                ap.get("doctor_name"),
                ap.get("status"),
                ap.get("payment_status"),
                ap.get("final_amount"),
            ]
            for ap in appointments
        ],
    )
    row += 1
    admissions = history.get("admissions") or []
    return _write_table(
        ws, row, "Admissions",
        ["Admitted", "Discharged", "Number", "Type", "Bed", "Doctor", "Status", "Reason"],
        [
            [
                adm.get("admission_date"),
                adm.get("discharge_date"),
                adm.get("admission_number"),
                adm.get("admission_type"),
                adm.get("bed_number"),
                adm.get("doctor_name"),
                adm.get("status"),
                adm.get("reason"),
            ]
            for adm in admissions
        ],
    )


def _write_consultations_section(ws, row: int, history: dict) -> int:
    rows = []
    for c in history.get("consultations") or []:
        diagnoses = "; ".join(
            d.get("diagnosis_name") or ""
            for d in (c.get("diagnoses") or [])
            if d.get("diagnosis_name")
        )
        vitals = c.get("vital_signs") or {}
        vitals_str = ", ".join(f"{k}={v}" for k, v in vitals.items() if v) if isinstance(vitals, dict) else ""
        rows.append([
            c.get("consultation_date"),
            c.get("consultation_number"),
            c.get("consultation_type"),
            c.get("doctor_name"),
            c.get("chief_complaint"),
            diagnoses,
            vitals_str,
            c.get("status"),
            c.get("consultation_fee"),
            c.get("follow_up_date"),
            c.get("notes"),
        ])
    return _write_table(
        ws, row, "Consultations",
        [
            "Date", "Number", "Type", "Doctor", "Chief complaint", "Diagnoses",
            "Vitals", "Status", "Fee", "Follow-up", "Notes",
        ],
        rows,
    )


def _write_prescriptions_section(ws, row: int, history: dict) -> int:
    rows = []
    for rx in history.get("prescriptions") or []:
        medicines = rx.get("medicines") or []
        if isinstance(medicines, str):
            try:
                medicines = json.loads(medicines)
            except Exception:
                medicines = []
        if not medicines:
            rows.append([
                rx.get("prescription_date"),
                rx.get("prescription_id"),
                rx.get("doctor_name"),
                rx.get("diagnosis"),
                "", "", "", "", "",
                rx.get("status"),
                rx.get("notes"),
            ])
            continue
        for m in medicines:
            if not isinstance(m, dict):
                continue
            rows.append([
                rx.get("prescription_date"),
                rx.get("prescription_id"),
                rx.get("doctor_name"),
                rx.get("diagnosis"),
                m.get("name") or m.get("medicine_name"),
                m.get("dosage"),
                m.get("frequency_schedule") or m.get("frequency"),
                m.get("duration"),
                m.get("instructions"),
                rx.get("status"),
                rx.get("notes"),
            ])
    return _write_table(
        ws, row, "Prescriptions",
        [
            "Date", "Prescription ID", "Doctor", "Diagnosis", "Medicine",
            "Dosage", "Frequency", "Duration", "Instructions", "Status", "Notes",
        ],
        rows,
    )


def _write_lab_section(ws, row: int, history: dict) -> int:
    rows = [
        [lo.get("order_date"), lo.get("test_name"), lo.get("status")]
        for lo in (history.get("lab_orders") or [])
    ]
    return _write_table(
        ws, row, "Lab orders",
        ["Date", "Test name", "Status"],
        rows,
    )


def _write_pharmacy_section(ws, row: int, history: dict) -> int:
    sales = history.get("pharmacy_sales") or []
    return _write_table(
        ws, row, "Pharmacy sales",
        ["Date", "Sale no.", "Doctor", "Payment", "Billing mode", "Status", "Amount"],
        [
            [
                s.get("sale_date"),
                s.get("sale_number"),
                s.get("doctor_name"),
                s.get("payment_type"),
                s.get("billing_mode"),
                s.get("status"),
                s.get("grand_total"),
            ]
            for s in sales
        ],
    )


def _write_billing_section(ws, row: int, history: dict) -> int:
    billing = history.get("billing") or {}
    row = _write_kv_block(ws, row, "Billing summary", [
        ("Total billed", billing.get("total_billed")),
        ("Total paid", billing.get("total_paid")),
        ("Outstanding", billing.get("outstanding")),
        ("Bill count", billing.get("bill_count")),
    ])
    row += 1
    bills = billing.get("bills") or []
    return _write_table(
        ws, row, "Bills",
        ["Date", "Bill no.", "Type", "Subtype", "Amount", "Paid", "Status"],
        [
            [
                b.get("bill_date"),
                b.get("bill_number"),
                b.get("bill_type"),
                b.get("bill_subtype"),
                b.get("total_amount"),
                b.get("amount_paid"),
                b.get("status"),
            ]
            for b in bills
        ],
    )


def _write_documents_section(ws, row: int, history: dict) -> int:
    docs = history.get("documents") or []
    return _write_table(
        ws, row, "Documents",
        ["Date", "Type", "Label", "Download path"],
        [
            [d.get("date"), d.get("type"), d.get("label"), d.get("download_url")]
            for d in docs
        ],
    )


def _write_kv_block(ws, start_row: int, title: str, pairs: Iterable[tuple[str, Any]]) -> int:
    ws.cell(row=start_row, column=1, value=title).font = _SECTION_FONT
    row = start_row + 1
    for label, value in pairs:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        cell = ws.cell(row=row, column=2, value=_cell_value(value))
        if isinstance(value, float):
            cell.number_format = "#,##0.00"
        row += 1
    return row


def _write_table(
    ws,
    start_row: int,
    title: str,
    headers: Sequence[str],
    data_rows: Sequence[Sequence[Any]],
) -> int:
    ws.cell(row=start_row, column=1, value=title).font = _SECTION_FONT
    row = start_row + 1
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    row += 1
    if not data_rows:
        ws.cell(row=row, column=1, value="No records")
        return row + 1
    for data in data_rows:
        for col, value in enumerate(data, start=1):
            cell = ws.cell(row=row, column=col, value=_cell_value(value))
            if isinstance(value, float):
                cell.number_format = "#,##0.00"
        row += 1
    return row


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _autosize(ws, max_width: int = 40) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        width = 10
        for cell in col_cells:
            if cell.value is None:
                continue
            width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
