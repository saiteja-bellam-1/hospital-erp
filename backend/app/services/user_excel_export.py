"""Excel roster export for Administration → Users.

This is a directory dump, not an import template: passwords are never written.
Hospital branding follows the same helper used by billing / EHR exports.
"""
from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.user import User
from app.utils.export_branding import apply_workbook_branding, hospital_brand_dict
from app.utils.time import format_system_dt, system_now

EXPORT_HEADERS = (
    "username",
    "email",
    "first_name",
    "last_name",
    "phone",
    "role",
    "additional_roles",
    "status",
    "specialization",
    "license_number",
    "qualification",
    "consultation_fee_inr",
    "inpatient_fee_inr",
    "inpatient_fee_charge_mode",
    "emergency_fee_inr",
    "experience_years",
    "created_at",
)

_COL_WIDTHS = (
    16, 28, 16, 16, 14, 18, 24, 12, 18, 16, 22, 16, 16, 22, 16, 16, 20,
)


def _primary_role(user: User) -> str:
    return user.role.name if user.role else ""


def _additional_roles(user: User) -> str:
    primary = _primary_role(user)
    extras = []
    seen = {primary} if primary else set()
    for role in user.roles or []:
        name = role.name or ""
        if name and name not in seen:
            seen.add(name)
            extras.append(name)
    return ";".join(extras)


def _cell(value) -> object:
    if value is None:
        return ""
    return value


def _user_row(user: User) -> list:
    return [
        user.username or "",
        user.email or "",
        user.first_name or "",
        user.last_name or "",
        user.phone or "",
        _primary_role(user),
        _additional_roles(user),
        "Active" if user.is_active else "Inactive",
        user.specialization or "",
        user.license_number or "",
        user.qualification or "",
        _cell(user.consultation_fee_inr),
        _cell(user.inpatient_fee_inr),
        getattr(user, "inpatient_fee_charge_mode", None) or "",
        _cell(user.emergency_fee_inr),
        _cell(user.experience_years),
        format_system_dt(user.created_at, empty=""),
    ]


def list_users_for_export(db: Session, *, include_super_admin: bool) -> list[User]:
    users = (
        db.query(User)
        .order_by(User.last_name, User.first_name, User.username)
        .all()
    )
    if include_super_admin:
        return users
    return [user for user in users if not user.has_role("super_admin")]


def build_users_xlsx(
    db: Session,
    *,
    include_super_admin: bool = True,
    hospital_id: Optional[int] = None,
) -> bytes:
    """Return .xlsx bytes for the current hospital's user roster."""
    users = list_users_for_export(db, include_super_admin=include_super_admin)
    hospital = hospital_brand_dict(db, hospital_id)
    active = sum(1 for user in users if user.is_active)
    inactive = len(users) - active
    exported_at = format_system_dt(system_now(), empty="")

    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    header_font = Font(bold=True, size=10, color="1A202C")
    header_fill = PatternFill("solid", fgColor="E8EEF5")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    meta_rows = [
        ("Report", "Users"),
        ("Exported at", exported_at),
        ("Total users", len(users)),
        ("Active", active),
        ("Inactive", inactive),
    ]
    header_row = apply_workbook_branding(ws, hospital, meta_rows)

    for col, title in enumerate(EXPORT_HEADERS, 1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(vertical="center")

    for offset, user in enumerate(users):
        row_idx = header_row + 1 + offset
        for col, value in enumerate(_user_row(user), 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin

    for i, width in enumerate(_COL_WIDTHS, 1):
        letter = get_column_letter(i)
        existing = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(existing, width)

    last_data_row = header_row + max(len(users), 0)
    ws.freeze_panes = f"A{header_row + 1}"
    last_col = get_column_letter(len(EXPORT_HEADERS))
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(header_row, last_data_row)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
